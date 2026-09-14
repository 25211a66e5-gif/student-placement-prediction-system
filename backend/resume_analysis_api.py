from flask import request, jsonify
from database import get_connection
from pathlib import Path
import json
import re
import tempfile
import os
import subprocess

# Resume analysis is intentionally format-flexible.
# The backend accepts any filename/extension and tries the appropriate extractor.
# Supported rich extraction paths:
# PDF  -> pypdf
# DOCX -> python-docx
# RTF  -> built-in RTF cleanup
# TXT/CSV/MD/HTML/XML/JSON -> text decoding
# Images -> OCR with Pillow + pytesseract when installed
# Legacy .DOC -> antiword/libreoffice when available
#
# There is no frontend extension whitelist.

MAX_FILE_SIZE = 10 * 1024 * 1024


ROLE_SKILLS = {
    "Software Developer": [
        "python", "java", "c++", "c", "javascript", "sql",
        "data structures", "algorithms", "git", "oops"
    ],
    "Full-Stack Developer": [
        "html", "css", "javascript", "react", "node", "sql",
        "python", "java", "rest api", "git", "mongodb"
    ],
    "Data Analyst": [
        "python", "sql", "excel", "power bi", "tableau",
        "statistics", "pandas", "numpy"
    ],
    "Data Scientist": [
        "python", "sql", "machine learning", "statistics",
        "pandas", "numpy", "scikit-learn", "data visualization"
    ],
    "ML / AI Engineer": [
        "python", "machine learning", "deep learning",
        "tensorflow", "pytorch", "scikit-learn", "numpy",
        "pandas", "sql"
    ],
    "Backend Developer": [
        "java", "python", "sql", "rest api", "spring boot",
        "flask", "django", "node", "git", "database"
    ],
    "Frontend Developer": [
        "html", "css", "javascript", "react", "typescript",
        "responsive", "git", "accessibility"
    ],
    "Embedded Engineer": [
        "c", "c++", "microcontroller", "embedded", "uart",
        "spi", "i2c", "rtos", "firmware"
    ],
    "Other Technical Role": [
        "python", "java", "c++", "sql", "git",
        "data structures", "algorithms"
    ],
}

COMMON_SKILLS = [
    "python", "java", "c", "c++", "c#", "javascript", "typescript",
    "html", "css", "sql", "mysql", "postgresql", "mongodb",
    "react", "angular", "vue", "node.js", "node", "spring", "spring boot",
    "django", "flask", "fastapi", "git", "github", "docker", "kubernetes",
    "aws", "azure", "gcp", "linux", "machine learning", "deep learning",
    "tensorflow", "pytorch", "scikit-learn", "pandas", "numpy",
    "power bi", "tableau", "excel", "data structures", "algorithms",
    "oops", "rest api", "microservices", "cybersecurity", "networking",
    "embedded", "microcontroller", "firmware", "uart", "spi", "i2c",
    "rtos", "statistics", "data visualization"
]

IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".bmp", ".tif", ".tiff", ".webp"}
TEXT_EXTENSIONS = {
    ".txt", ".text", ".md", ".csv", ".tsv", ".log", ".json", ".xml",
    ".html", ".htm", ".rtf"
}


def ensure_resume_tables():
    connection = get_connection()
    try:
        connection.execute("""
            CREATE TABLE IF NOT EXISTS resume_profiles (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                user_id INTEGER NOT NULL UNIQUE,
                filename TEXT NOT NULL,
                file_type TEXT NOT NULL,
                resume_text TEXT NOT NULL,
                detected_skills TEXT NOT NULL DEFAULT '[]',
                projects_count INTEGER NOT NULL DEFAULT 0,
                internships_count INTEGER NOT NULL DEFAULT 0,
                certifications_count INTEGER NOT NULL DEFAULT 0,
                education_score REAL NOT NULL DEFAULT 0,
                role_match_score REAL NOT NULL DEFAULT 0,
                resume_score REAL NOT NULL DEFAULT 0,
                placement_estimate REAL NOT NULL DEFAULT 0,
                missing_skills TEXT NOT NULL DEFAULT '[]',
                strengths TEXT NOT NULL DEFAULT '[]',
                recommendations TEXT NOT NULL DEFAULT '[]',
                created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
                updated_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
                FOREIGN KEY (user_id) REFERENCES users(id) ON DELETE CASCADE
            )
        """)
        connection.commit()
    finally:
        connection.close()


def _extract_pdf(path):
    try:
        from pypdf import PdfReader
        reader = PdfReader(path)
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    except ImportError:
        try:
            import fitz
            doc = fitz.open(path)
            return "\n".join(page.get_text() for page in doc)
        except ImportError:
            raise RuntimeError(
                "PDF reader is not installed. Run: pip install pypdf"
            )


def _extract_docx(path):
    try:
        from docx import Document
    except ImportError:
        raise RuntimeError(
            "DOCX reader is not installed. Run: pip install python-docx"
        )
    doc = Document(path)
    chunks = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            chunks.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(chunks)


def _extract_rtf(path):
    raw = Path(path).read_text(encoding="utf-8", errors="ignore")
    raw = re.sub(r"\\'[0-9a-fA-F]{2}", " ", raw)
    raw = re.sub(r"\\[a-zA-Z]+\d* ?", " ", raw)
    raw = raw.replace("{", " ").replace("}", " ")
    return raw


def _extract_image(path):
    try:
        from PIL import Image
        import pytesseract
    except ImportError:
        raise RuntimeError(
            "Image resume OCR requires Pillow and pytesseract. "
            "Run: pip install Pillow pytesseract. "
            "You also need the Tesseract OCR application installed."
        )
    image = Image.open(path)
    return pytesseract.image_to_string(image)


def _extract_legacy_doc(path):
    # Try antiword first, then LibreOffice/soffice.
    try:
        result = subprocess.run(
            ["antiword", path],
            capture_output=True, text=True, timeout=20
        )
        if result.returncode == 0 and result.stdout.strip():
            return result.stdout
    except (FileNotFoundError, subprocess.TimeoutExpired):
        pass

    try:
        with tempfile.TemporaryDirectory() as folder:
            result = subprocess.run(
                [
                    "libreoffice", "--headless", "--convert-to", "txt:Text",
                    "--outdir", folder, path
                ],
                capture_output=True, text=True, timeout=30
            )
            txt = Path(folder) / (Path(path).stem + ".txt")
            if result.returncode == 0 and txt.exists():
                return txt.read_text(encoding="utf-8", errors="ignore")
    except (FileNotFoundError, subprocess.TimeoutExpired):
        pass

    raise RuntimeError(
        "This legacy .DOC file needs a document converter. "
        "Install LibreOffice or antiword on the machine running Flask."
    )


def _extract_text_file(path):
    raw = Path(path).read_bytes()
    for encoding in ("utf-8", "utf-16", "cp1252", "latin-1"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="ignore")


def extract_resume_text(path, extension):
    extension = extension.lower()

    if extension == ".pdf":
        return _extract_pdf(path)
    if extension == ".docx":
        return _extract_docx(path)
    if extension == ".doc":
        return _extract_legacy_doc(path)
    if extension == ".rtf":
        return _extract_rtf(path)
    if extension in IMAGE_EXTENSIONS:
        return _extract_image(path)
    if extension in TEXT_EXTENSIONS:
        return _extract_text_file(path)

    # No extension restriction: try text decoding first. This handles
    # resumes exported with unusual extensions.
    text = _extract_text_file(path)
    printable = sum(1 for c in text if c.isprintable() or c in "\n\r\t")
    ratio = printable / max(1, len(text))
    if ratio >= 0.85:
        return text

    raise RuntimeError(
        "The uploaded file could not be converted into readable resume text. "
        "Try a text-based document, PDF/DOCX, or an image with readable text."
    )


def clean_text(text):
    text = text.replace("\x00", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def contains_skill(text, skill):
    skill = skill.lower()
    if skill in {"c", "c++", "c#"}:
        return re.search(r"(?<![a-z0-9+#])" + re.escape(skill) + r"(?![a-z0-9+#])", text) is not None
    if " " in skill or "/" in skill or "+" in skill or "." in skill:
        return skill in text
    return re.search(r"(?<![a-z0-9])" + re.escape(skill) + r"(?![a-z0-9])", text) is not None


def _count_evidence(text, patterns, cap):
    return min(cap, max(0, sum(len(re.findall(p, text)) for p in patterns)))


def analyze_resume(text, target_role, profile):
    lower = text.lower()
    role = target_role or "Software Developer"
    required = ROLE_SKILLS.get(role, ROLE_SKILLS["Other Technical Role"])

    detected_role_skills = [s for s in required if contains_skill(lower, s)]
    detected_common = [s for s in COMMON_SKILLS if contains_skill(lower, s)]
    detected = list(dict.fromkeys(detected_role_skills + detected_common))

    projects_count = _count_evidence(
        lower,
        [r"\bprojects?\b", r"\bdeveloped\b", r"\bimplemented\b", r"\bbuilt\b"],
        10
    )
    internships_count = _count_evidence(
        lower,
        [r"\binternships?\b", r"\bintern\b", r"\binterned\b", r"\btrainee\b"],
        5
    )
    certifications_count = _count_evidence(
        lower,
        [r"\bcertifications?\b", r"\bcertified\b", r"\bcertificate\b"],
        10
    )

    education_terms = [
        "b.tech", "b.e.", "bachelor", "engineering", "cgpa",
        "university", "college", "education", "degree"
    ]
    education_hits = sum(term in lower for term in education_terms)
    education_score = min(100, education_hits / len(education_terms) * 100)

    role_match = round(
        len(detected_role_skills) / max(1, len(required)) * 100, 2
    )

    evidence_score = min(
        100,
        projects_count * 8 +
        internships_count * 10 +
        certifications_count * 4
    )

    completeness_terms = [
        "education", "skills", "project", "experience",
        "certification", "contact", "objective"
    ]
    completeness = sum(term in lower for term in completeness_terms)
    completeness_score = min(
        100, completeness / len(completeness_terms) * 100
    )

    resume_score = round(
        role_match * 0.45 +
        evidence_score * 0.25 +
        education_score * 0.10 +
        completeness_score * 0.20,
        2
    )

    missing = [s for s in required if s not in detected_role_skills]

    strengths = []
    if role_match >= 70:
        strengths.append("Strong technical alignment with the selected role.")
    elif role_match >= 40:
        strengths.append("Moderate technical alignment with the selected role.")
    else:
        strengths.append("The resume needs stronger role-specific technical evidence.")

    if projects_count:
        strengths.append(f"Project evidence detected ({projects_count} signal(s)).")
    if internships_count:
        strengths.append(f"Internship/experience evidence detected ({internships_count} signal(s)).")
    if certifications_count:
        strengths.append(f"Certification evidence detected ({certifications_count} signal(s)).")

    recommendations = []
    if missing:
        recommendations.append(
            "Build or document projects using: " + ", ".join(missing[:6]) + "."
        )
    if projects_count == 0:
        recommendations.append("Add at least one detailed technical project with technologies and outcomes.")
    if internships_count == 0:
        recommendations.append("Add internship, training, freelance, or practical experience if available.")
    if certifications_count == 0:
        recommendations.append("Add relevant certifications or verified learning achievements.")
    if completeness_score < 70:
        recommendations.append("Improve resume completeness with clear education, skills, projects and experience sections.")

    academic_score = 0
    if profile:
        cgpa = float(profile["cgpa"] or 0)
        tenth = float(profile["tenth_percentage"] or 0)
        twelfth = float(profile["twelfth_percentage"] or 0)
        backlogs = int(profile["backlogs"] or 0)
        academic_score = (
            min(100, cgpa / 10 * 100) * 0.50 +
            tenth * 0.20 +
            twelfth * 0.20 +
            max(0, 100 - backlogs * 20) * 0.10
        )

    placement_estimate = round(
        academic_score * 0.45 +
        role_match * 0.25 +
        resume_score * 0.30,
        2
    )

    return {
        "detected_skills": detected,
        "projects_count": projects_count,
        "internships_count": internships_count,
        "certifications_count": certifications_count,
        "education_score": round(education_score, 2),
        "role_match_score": role_match,
        "resume_score": resume_score,
        "placement_estimate": placement_estimate,
        "missing_skills": missing,
        "strengths": strengths,
        "recommendations": recommendations,
    }


def serialize(row):
    return {
        "filename": row["filename"],
        "file_type": row["file_type"],
        "detected_skills": json.loads(row["detected_skills"] or "[]"),
        "projects_count": row["projects_count"],
        "internships_count": row["internships_count"],
        "certifications_count": row["certifications_count"],
        "education_score": row["education_score"],
        "role_match_score": row["role_match_score"],
        "resume_score": row["resume_score"],
        "placement_estimate": row["placement_estimate"],
        "missing_skills": json.loads(row["missing_skills"] or "[]"),
        "strengths": json.loads(row["strengths"] or "[]"),
        "recommendations": json.loads(row["recommendations"] or "[]"),
    }


def register_resume_routes(app):
    ensure_resume_tables()

    @app.post("/api/resume/upload")
    def upload_resume():
        user_id = request.form.get("user_id")
        target_role = request.form.get("target_job_role")
        uploaded = request.files.get("resume")

        if not user_id:
            return jsonify({"error": "User ID is required."}), 400
        if uploaded is None or not uploaded.filename:
            return jsonify({"error": "Please select a resume file."}), 400

        connection = get_connection()
        try:
            profile = connection.execute("""
                SELECT *
                FROM student_profiles
                WHERE user_id = ?
            """, (user_id,)).fetchone()

            if profile is None:
                return jsonify({
                    "error": "Please complete your student profile first."
                }), 400

            role = target_role or profile["target_job_role"]
            original_name = Path(uploaded.filename).name
            extension = Path(original_name).suffix.lower()

            raw = uploaded.read(MAX_FILE_SIZE + 1)
            if len(raw) > MAX_FILE_SIZE:
                return jsonify({
                    "error": "Resume file must be 10 MB or smaller."
                }), 400

            with tempfile.NamedTemporaryFile(
                suffix=extension, delete=False
            ) as temp:
                temp.write(raw)
                temp_path = temp.name

            try:
                text = clean_text(extract_resume_text(temp_path, extension))
            finally:
                try:
                    os.unlink(temp_path)
                except OSError:
                    pass

            if len(text) < 80:
                return jsonify({
                    "error": (
                        "The file was uploaded, but not enough readable resume "
                        "text could be extracted. If it is a scanned/image resume, "
                        "enable OCR support on the backend."
                    )
                }), 400

            analysis = analyze_resume(text, role, profile)

            connection.execute("""
                INSERT INTO resume_profiles (
                    user_id, filename, file_type, resume_text,
                    detected_skills, projects_count, internships_count,
                    certifications_count, education_score, role_match_score,
                    resume_score, placement_estimate, missing_skills,
                    strengths, recommendations, updated_at
                )
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, CURRENT_TIMESTAMP)
                ON CONFLICT(user_id) DO UPDATE SET
                    filename=excluded.filename,
                    file_type=excluded.file_type,
                    resume_text=excluded.resume_text,
                    detected_skills=excluded.detected_skills,
                    projects_count=excluded.projects_count,
                    internships_count=excluded.internships_count,
                    certifications_count=excluded.certifications_count,
                    education_score=excluded.education_score,
                    role_match_score=excluded.role_match_score,
                    resume_score=excluded.resume_score,
                    placement_estimate=excluded.placement_estimate,
                    missing_skills=excluded.missing_skills,
                    strengths=excluded.strengths,
                    recommendations=excluded.recommendations,
                    updated_at=CURRENT_TIMESTAMP
            """, (
                user_id, original_name, extension.lstrip(".") or "unknown", text,
                json.dumps(analysis["detected_skills"]),
                analysis["projects_count"],
                analysis["internships_count"],
                analysis["certifications_count"],
                analysis["education_score"],
                analysis["role_match_score"],
                analysis["resume_score"],
                analysis["placement_estimate"],
                json.dumps(analysis["missing_skills"]),
                json.dumps(analysis["strengths"]),
                json.dumps(analysis["recommendations"]),
            ))
            connection.commit()

            row = connection.execute("""
                SELECT * FROM resume_profiles WHERE user_id = ?
            """, (user_id,)).fetchone()

            return jsonify({
                "message": "Resume uploaded and analyzed successfully.",
                "analysis": serialize(row),
                "disclaimer": (
                    "This preliminary resume estimate is not a guarantee of employment. "
                    "The final project result should combine the student's academic profile, "
                    "assessment, coding test, interview and role-specific evidence."
                )
            }), 200

        except Exception as error:
            connection.rollback()
            return jsonify({
                "error": "Resume analysis failed.",
                "message": str(error)
            }), 500
        finally:
            connection.close()

    @app.get("/api/resume/<int:user_id>")
    def get_resume_analysis(user_id):
        connection = get_connection()
        try:
            row = connection.execute("""
                SELECT * FROM resume_profiles WHERE user_id = ?
            """, (user_id,)).fetchone()

            if row is None:
                return jsonify({"error": "No resume has been uploaded yet."}), 404

            return jsonify({"analysis": serialize(row)}), 200
        finally:
            connection.close()
